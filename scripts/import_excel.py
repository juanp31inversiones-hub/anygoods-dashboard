import openpyxl, datetime, sys, warnings
warnings.filterwarnings('ignore')
UP = sys.argv[1]
OUT = sys.argv[2]

def q(v):
    if v is None: return 'null'
    if isinstance(v, (int, float)): return repr(round(v, 4))
    return "'" + str(v).replace("'", "''") + "'"

# Catálogo maestro: clave = descripción en Purchases -> (descripción, upc, asin, retail Amazon, brand)
P = {
 'ARM & HAMMER SpinBrush Pro Extra Soft 1ct': ('ARM & HAMMER SpinBrush Pro Extra Soft 1ct','766878000787','B001J4ID5K',11.98,'Arm & Hammer'),
 'BASIS Soap 4oz sensitive Skin Bar': ('BASIS Soap 4oz Sensitive Skin Bar','72140857004','B008CYC10I',19.77,'Basis'),
 'BIGEN Permanent Powder Hair Color 59 Oriental Black 1 ea (Pack of 6)': ('BIGEN Permanent Powder Hair Color #59 Oriental Black','33859905592','B012JAQGOM',59.63,'Bigen'),
 'CHERIE SWEET 16 oz organic matcha powder 18 bags/case': ('CHERIE SWEET 16oz Organic Matcha Powder',None,'B07P6454Z1',16.37,'Cherie Sweet'),
 'CLEANING HEROES 99% Isopropyl Alcohol Spray, 32 Fl. Oz': ('CLEANING HEROES 99% Isopropyl Alcohol Spray 32oz',None,None,None,'Cleaning Heroes'),
 'CLEANING HEROES Oven & Grill 32 Oz': ('CLEANING HEROES Oven & Grill 32oz',None,None,None,'Cleaning Heroes'),
 'DOVE SS 24/1.7oz CLINICAL PRO+CARE COOL ESSENTIAL': ('DOVE Clinical Pro+Care Cool Essential 1.7oz','79400008787','B001QVNPSQ',8.53,'Dove'),
 'ECOCO Black Castor & Flax Seed Oil 32oz': ('ECOCO Black Castor & Flax Seed Oil Gel 32oz','748378004229','B01NB20C2V',12.90,'Ecoco'),
 'ECOCO COCONUT STYLING GEL 6/32oz': ('ECOCO Coconut Styling Gel 32oz','748378004342','B075M9PJ4D',14.99,'Ecoco'),
 'ECOCO Curl & Wave Styling Gel Pink 32oz': ('ECOCO Curl & Wave Styling Gel Pink 32oz','748378001013','B001T8MJC4',12.95,'Ecoco'),
 'ECOCO Krystal Styling Gel Clear 32oz': ('ECOCO Krystal Styling Gel Clear 32oz','748378001150','B08Z3MT5WM',17.88,'Ecoco'),
 'ECOCO Olive Oil Styling Gel Green 32oz': ('ECOCO Olive Oil Styling Gel Green 32oz','748378001594','B003E7UNE4',14.66,'Ecoco'),
 'ECOCO Olive Oil Styling Gel Green 80oz': ('ECOCO Olive Oil Styling Gel Green 80oz','748378001402','B003E82NG4',21.49,'Ecoco'),
 'ECOCO Sport Styling Gel Blue 16oz': ('ECOCO Sport Styling Gel Blue 16oz','748378001235','B005I4PJ5Y',8.54,'Ecoco'),
 'ELASTA Glaze 6oz': ('ELASTA QP Glaze Mango Butter 6oz','802535554061','B00CWFRQUI',7.98,'Elasta QP'),
 'EQUAL packets 50ct': ('EQUAL Packets 50ct','25810950','B0005XMU00',9.49,'Equal'),
 'HI & DRILL Roll On 7oz': ('HI & DRY Roll On Powder Fresh Pink 1.7oz','309975122039','B00N21HG7E',14.72,'Hi & Dry'),
 'JUST FOR ME Relaxer No Lye Conditioning': ('JUST FOR ME Relaxer No Lye Conditioning Creme','802535373006','B0755B2WZ2',13.24,'Just For Me'),
 'LA LOOKS Styling Gel Mega Shaper Hold (lv 9)20oz': ('LA LOOKS Styling Gel Mega Shaper Hold Lv9 20oz','41670132356','B005JT5A84',11.46,'LA Looks'),
 'NUGGET Shoe Polish 2oz - Black': ('NUGGET Shoe Polish 2oz Black','7501058743121','B001DVWUO4',10.90,'Nugget'),
 'OLD SPICE Body Wash, Fiji With Palm Tree - 16 Oz - 4 Pack': ('OLD SPICE Body Wash Fiji 16oz 4-Pack','37000267904','B00498HG0G',7.98,'Old Spice'),
}

sql = ['begin;']
upp = {}
for k, (d, upc, asin, retail, brand) in P.items():
    sql.append(f"insert into products (description, upc, asin, retail_price, brand) values ({q(d)},{q(upc)},{q(asin)},{q(retail)},{q(brand)});")

# Contrapartes
CP = {'Contar Market':'supplier','Estmed Brands Inc':'supplier','Cherie Sweet':'supplier','Cleaning Heroes':'supplier',
      'EC Access':'supplier','Amazon':'marketplace','Ebay':'marketplace','Keepa':'service','AzInsights':'service',
      'Google':'service','SmartScout':'service','Carbon 6':'service','Stella Consulting':'service'}
ALIAS = {'Gmail':'Google','Stella Cosulting':'Stella Consulting'}
for n, k in CP.items():
    sql.append(f"insert into counterparties (name, kind) values ({q(n)},'{k}');")

# ===== Compras =====
wb = openpyxl.load_workbook(f'{UP}/2f66fcde-Product_Inventory_ANY_GOODS_3.xlsx', data_only=True)
pos = {}
for r in wb['Purchases'].iter_rows(min_row=4, max_row=40, values_only=True):
    r = list(r) + [None]*30
    if not (r[5] and r[6] and isinstance(r[1], datetime.datetime)): continue   # omite fila sin fecha (cotización)
    inv = str(r[4]).replace('.0','')
    key = (r[5], inv)
    po = pos.setdefault(key, {'date': r[1].date(), 'arr': r[2].date() if isinstance(r[2], datetime.datetime) else None,
                              'sup': 0.0, 'amz': 0.0, 'lines': []})
    units = float(r[13]); cost = float(r[15])
    sup_t = float(r[16]) if isinstance(r[16], (int, float)) else 0.0
    final = float(r[23]) if isinstance(r[23], (int, float)) else cost/units
    extra = max(0.0, final*units - cost)           # transporte proveedor + Amazon según la hoja
    po['sup'] += sup_t
    po['amz'] += max(0.0, extra - sup_t)
    po['lines'].append((P[' '.join(r[6].split())][0], float(r[12]) if r[12] else None, int(units), cost/units, extra, final))

for (sup, inv), po in pos.items():
    sql.append(f"""insert into purchase_orders (supplier_id, invoice_number, purchase_date, arrival_date, status, supplier_shipping, amazon_shipping)
  select id, {q(inv)}, {q(str(po['date']))}, {q(str(po['arr']) if po['arr'] else None)}, 'sent_to_amazon', {round(po['sup'],2)}, {round(po['amz'],2)} from counterparties where name={q(sup)};""")
    for d, cases, units, uc, extra, final in po['lines']:
        sql.append(f"""with l as (insert into purchase_order_lines (purchase_order_id, product_id, cases, units, unit_cost, allocated_costs, landed_unit_cost)
  select po.id, p.id, {q(cases)}, {units}, {q(uc)}, {round(extra,2)}, {q(final)}
  from purchase_orders po join counterparties c on c.id=po.supplier_id, products p
  where c.name={q(sup)} and po.invoice_number={q(inv)} and p.description={q(d)} returning id, product_id)
insert into inventory_movements (product_id, po_line_id, movement, qty, occurred_on, channel, reference)
  select product_id, id, 'purchase', {units}, {q(str(po['arr'] or po['date']))}, 'amazon', {q('INV '+inv)} from l;""")

# ===== Banco =====
wb = openpyxl.load_workbook(f'{UP}/2f9518aa-ANY_GOODS_ADMIN.xlsx', data_only=True)
EG = {'Herramientas':'software','Legal':'legal','Envios':'shipping','Inventario':'inventory','Devolucion':'fees'}
n = 0
bank = []
for r in wb['I+G 2025'].iter_rows(values_only=True):
    r = list(r) + [None]*20
    # Ingresos
    if isinstance(r[1], datetime.datetime) and isinstance(r[8], (int, float)):
        n += 1
        desc, cat, payer, amt = r[4], r[5], r[7], r[8]
        if cat == 'Aporte':
            c, partner, cp = 'capital_contribution', payer, None
        else:
            partner = None
            cp = 'Amazon' if (payer == 'Amazon' or 'Amazon' in desc) else payer
            c = 'amazon_payout' if cp == 'Amazon' else 'other_sales'
        cp = cp if cp in CP else None
        bank.append(f"({q(str(r[1].date()))},{q(desc)},{round(amt,2)},'{c}',{q(cp)},{q(partner)},'xls-ig-{n}',null)")
    # Egresos
    if isinstance(r[10], datetime.datetime) and isinstance(r[17], (int, float)) and r[14]:
        n += 1
        cp = ALIAS.get(r[15], r[15])
        cp = cp if cp in CP else None
        note = None
        if r[14] == 'Devolucion': note = 'Excel: Devolucion. Cargo de Amazon/eBay a la cuenta (saldo negativo del seller).'
        if 'Initiation Fee' in (r[13] or ''): note = 'Excel lo clasifica como Inventario. Revisar: parece fee de programa, no mercancía.'
        bank.append(f"({q(str(r[10].date()))},{q(r[13])},{-round(r[17],2)},'{EG[r[14]]}',{q(cp)},null,'xls-ig-{n}',{q(note)})")

sql.append("""insert into bank_transactions (bank_account_id, txn_date, description, amount, category, counterparty_id, partner_id, external_id, notes)
select (select id from bank_accounts where name='Mercury'), v.d::date, v.descr, v.amt, v.cat::txn_category, c.id, p.id, v.ext, v.note
from (values """ + ",\n".join(bank) + """) v(d,descr,amt,cat,cp,partner,ext,note)
left join counterparties c on c.name = v.cp left join partners p on p.short_name = v.partner;""")
# Concilia pagos al proveedor con la PO de mismo proveedor y mismo valor de producto
sql.append("""update bank_transactions t set purchase_order_id = x.po_id from (
  select distinct on (t.id) t.id txn_id, po.id po_id
  from bank_transactions t join purchase_orders po on po.supplier_id = t.counterparty_id
  join lateral (select sum(units*unit_cost) v from purchase_order_lines where purchase_order_id = po.id) s on true
  where t.category='inventory' and abs(-t.amount - s.v) < 0.01 and abs(t.txn_date - po.purchase_date) <= 45
  order by t.id, abs(t.txn_date - po.purchase_date)) x where t.id = x.txn_id;""")
sql.append('commit;')
open(OUT, 'w').write('\n'.join(sql))
print(len(pos), 'POs,', sum(len(p['lines']) for p in pos.values()), 'lines,', n, 'bank txns')
